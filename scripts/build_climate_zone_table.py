"""Build config/climate_zones.yaml from the official NatHERS postcode file.

The authoritative NatHERS "Climate Zones by postcode" file (Excel or PDF) maps
each Australian postcode to a Primary (and sometimes Secondary/Tertiary) NatHERS
climate zone. This script converts it into PlanAssess's climate_zones.yaml:

* a postcode with a single zone  -> high-confidence `postcodes` entry;
* a postcode spanning >1 zone     -> low-confidence `ambiguous` entry (so the
  review gate flags it instead of silently picking one).

It is NOT run automatically and the official file is NOT vendored (licensing +
size). Download it from the source below, then run:

    pip install openpyxl pdfplumber   # xlsx and/or pdf support
    python scripts/build_climate_zone_table.py NatHERSclimatezonesSep2025.xlsx \
        --out config/climate_zones.yaml

Source (verify currency before use):
    https://www.nathers.gov.au/climate-zone-postcodes
"""

from __future__ import annotations

import argparse
import datetime as _dt
from collections import OrderedDict
from pathlib import Path


def _rows_from_xlsx(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]


def _rows_from_pdf(path: Path) -> list[list[str]]:
    import pdfplumber

    rows: list[list[str]] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for row in table:
                    rows.append(["" if c is None else str(c).strip() for c in row])
    return rows


def _find_columns(header: list[str]) -> tuple[int, list[int]]:
    """Return (postcode_col, [zone_cols]) by matching header names."""
    pc_col, zone_cols = None, []
    for i, name in enumerate(header):
        low = name.lower()
        if pc_col is None and "postcode" in low:
            pc_col = i
        if any(k in low for k in ("primary", "secondary", "tertiary", "climate zone", "zone")):
            zone_cols.append(i)
    if pc_col is None:
        pc_col = 0
    if not zone_cols:
        zone_cols = [i for i in range(len(header)) if i != pc_col]
    return pc_col, zone_cols


def build(rows: list[list[str]]) -> tuple[dict, dict]:
    if not rows:
        raise SystemExit("No rows parsed from the input file.")
    header = rows[0]
    pc_col, zone_cols = _find_columns(header)

    postcodes: "OrderedDict[str, dict]" = OrderedDict()
    ambiguous: "OrderedDict[str, dict]" = OrderedDict()
    for row in rows[1:]:
        if pc_col >= len(row):
            continue
        pc = "".join(ch for ch in row[pc_col] if ch.isdigit())
        if len(pc) != 4:
            continue
        zones: list[int] = []
        for c in zone_cols:
            if c < len(row) and row[c].strip().isdigit():
                z = int(row[c].strip())
                if 1 <= z <= 69 and z not in zones:
                    zones.append(z)
        if not zones:
            continue
        if len(zones) == 1:
            postcodes[pc] = {"zone": zones[0], "confidence": 0.9}
        else:
            ambiguous[pc] = {
                "zones": zones,
                "confidence": 0.4,
                "note": "Postcode spans multiple NatHERS zones (primary/secondary/tertiary).",
            }
    return postcodes, ambiguous


def to_yaml(postcodes: dict, ambiguous: dict, source: str) -> str:
    import yaml

    doc = {
        "version": f"nathers-{_dt.date.today().isoformat()}",
        "source": source,
        "postcodes": {k: dict(v) for k, v in postcodes.items()},
        "ambiguous": {k: dict(v) for k, v in ambiguous.items()},
    }
    return yaml.safe_dump(doc, sort_keys=False, default_flow_style=False)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input", type=Path, help="Official NatHERS postcode file (.xlsx or .pdf)")
    ap.add_argument("--out", type=Path, default=Path("config/climate_zones.yaml"))
    args = ap.parse_args()

    suffix = args.input.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        rows = _rows_from_xlsx(args.input)
    elif suffix == ".pdf":
        rows = _rows_from_pdf(args.input)
    else:
        raise SystemExit(f"Unsupported input '{suffix}'. Provide .xlsx or .pdf.")

    postcodes, ambiguous = build(rows)
    args.out.write_text(to_yaml(postcodes, ambiguous, str(args.input.name)), encoding="utf-8")
    print(f"wrote {args.out}: {len(postcodes)} postcodes, {len(ambiguous)} ambiguous")


if __name__ == "__main__":
    main()
